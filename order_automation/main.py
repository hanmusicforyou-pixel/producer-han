"""영암부부농원 발주 자동화 - 실행 진입점

사용법:
  python main.py run                 # 지금 시각 기준으로 할 일(수집/회신확인)을 한 번 수행  ← GitHub Actions 가 호출
  python main.py run --date 2026-10-05 --force
  python main.py collect [--dry-run] # 1단계만
  python main.py check-reply         # 2단계만
  python main.py daemon              # PC 에서 계속 켜두는 스케줄러 (매일 10:00 수집, 이후 30분마다 회신 확인)
  python main.py announce --to a@x.com,b@y.com   # 업체에 접수 안내문 발송 (최초 1회)
  python main.py preview-announcement            # 안내문 내용만 출력
  python main.py test-login          # 네이버 IMAP/SMTP 로그인 확인
  python main.py demo [--sample 파일.xlsx]        # 네이버 접속 없이 전체 흐름 시연 (data/demo 에 결과 저장)
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from config import KST, Config, load_config


def setup_logging(cfg: Config) -> None:
    cfg.data_dir.mkdir(parents=True, exist_ok=True)
    log_dir = cfg.data_dir / "logs"
    log_dir.mkdir(exist_ok=True)
    handlers = [logging.StreamHandler(sys.stdout),
                logging.FileHandler(log_dir / f"{datetime.now(KST):%Y%m%d}.log", encoding="utf-8")]
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s", handlers=handlers)


def real_mail(cfg: Config, dry_run: bool):
    from naver_mail import NaverMail
    problems = cfg.validate()
    if problems:
        for p in problems:
            print("설정 오류:", p)
        print("→ order_automation/.env (PC) 또는 GitHub Secrets 에 값을 넣어주세요. 예시는 .env.example 참고.")
        sys.exit(2)
    return NaverMail(cfg.naver_id, cfg.naver_password, display_name=cfg.my_name, dry_run=dry_run,
                     outbox_dir=cfg.data_dir / "outbox")


def parse_date(s: str | None) -> date | None:
    if not s:
        return None
    if "-" in s:
        return date.fromisoformat(s)
    m, d = s.replace("/", ".").split(".")[:2]  # 10.5 또는 10/5
    today = datetime.now(KST).date()
    return date(today.year, int(m), int(d))


def print_result(res: dict) -> None:
    import json
    print(json.dumps(res, ensure_ascii=False, indent=2, default=str))


def cmd_run(cfg, args):
    from pipeline import Pipeline
    with real_mail(cfg, args.dry_run) as mail:
        print_result(Pipeline(cfg, mail, args.dry_run).run(parse_date(args.date), force=args.force))


def cmd_collect(cfg, args):
    from pipeline import Pipeline
    with real_mail(cfg, args.dry_run) as mail:
        p = Pipeline(cfg, mail, args.dry_run)
        target = parse_date(args.date) or p.today()
        print_result(p.collect(target, force=args.force))


def cmd_check_reply(cfg, args):
    from pipeline import Pipeline
    with real_mail(cfg, args.dry_run) as mail:
        p = Pipeline(cfg, mail, args.dry_run)
        print_result(p.check_reply(parse_date(args.date) or p.today()))


def cmd_daemon(cfg, args):
    import time as _time
    import schedule
    from pipeline import Pipeline

    def job():
        try:
            with real_mail(cfg, args.dry_run) as mail:
                res = Pipeline(cfg, mail, args.dry_run).run()
                logging.getLogger("daemon").info("run 결과: %s", {k: (v.get("status") if isinstance(v, dict) else v) for k, v in res.items()})
        except Exception:  # noqa: BLE001
            logging.getLogger("daemon").exception("실행 중 오류 (다음 주기에 재시도)")

    ct = cfg.collect_time
    schedule.every().day.at(f"{ct.hour:02d}:{ct.minute:02d}").do(job)
    schedule.every(cfg.reply_check_interval_min).minutes.do(job)
    print(f"스케줄러 시작: 매일 {ct.hour:02d}:{ct.minute:02d} 수집, {cfg.reply_check_interval_min}분마다 상태 확인. "
          f"운영 종료일 {cfg.end_date}. (Ctrl+C 로 종료)")
    job()
    while True:
        schedule.run_pending()
        _time.sleep(20)


def cmd_announce(cfg, args):
    from pipeline import Pipeline
    recipients = [x.strip() for x in (args.to or "").split(",") if x.strip()]
    if not recipients:
        print("--to a@x.com,b@y.com 형식으로 받는 업체 메일을 지정하세요.")
        sys.exit(2)
    with real_mail(cfg, args.dry_run) as mail:
        for r in Pipeline(cfg, mail, args.dry_run).announce(recipients):
            print(("DRY-RUN " if r.dry_run else "발송 ") + ", ".join(r.to) + " | " + r.subject)


def cmd_preview_announcement(cfg, args):
    from fake_mail import FakeMail
    from pipeline import Pipeline
    print(Pipeline(cfg, FakeMail(cfg.my_address), True).render_announcement())


def cmd_test_login(cfg, args):
    import smtplib
    with real_mail(cfg, False) as mail:
        hdrs = mail.list_headers(datetime.now(KST).date() - timedelta(days=1))
        print(f"IMAP OK - 최근 2일 받은메일 {len(hdrs)}건")
        for h in hdrs[-5:]:
            print("  ", h.short())
    with smtplib.SMTP_SSL(mail.SMTP_HOST, mail.SMTP_PORT, timeout=30) as smtp:
        try:
            smtp.login(mail.login_id, mail.password)
        except smtplib.SMTPAuthenticationError:
            smtp.login(mail.address, mail.password)
    print("SMTP OK - 발송 로그인 성공")


def cmd_demo(cfg, args):
    """네이버 접속 없이 전체 흐름 시연. 샘플 엑셀을 두 업체가 보낸 것으로 가정."""
    import io
    import openpyxl
    import order_excel as oe
    from fake_mail import FakeMail
    from pipeline import Pipeline, MARK_A, MARK_B

    demo_dir = cfg.data_dir / "demo"
    cfg.data_dir = demo_dir
    mail = FakeMail(cfg.my_address, outbox_dir=demo_dir / "outbox")
    mail.display_name = cfg.my_name
    p = Pipeline(cfg, mail, dry_run=True)
    target = parse_date(args.date) or p.today()

    if args.sample:
        rows = oe.parse_order_bytes(Path(args.sample).read_bytes(), Path(args.sample).name)
    else:
        rows = _synthetic_rows()
    half = max(1, len(rows) * 2 // 3)
    vendors = [("두리농", "doorinong@example.com", rows[:half]), ("햇살농장", "sunfarm@example.com", rows[half:])]
    when = datetime.combine(target, cfg.collect_time, KST) - timedelta(minutes=40)
    for name, addr, vrows in vendors:
        for r in vrows:
            r.values["보내는분성명"] = name
        data = oe.build_consolidated_xlsx(vrows, target)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        for extra in ("품목요약", "업체별요약"):
            del wb[extra]
        buf = io.BytesIO(); wb.save(buf)
        mail.inject(name, addr, f"{target.month}월 {target.day}일 {cfg.subject_keyword}", f"{name} 발주서 보냅니다.", when,
                    [(f"{name}_발주서.xlsx", buf.getvalue())])
        when += timedelta(minutes=7)

    print("=" * 70, "\n[1단계] 업체 발주서 수집 → 통합발주 → 박경원님 발송 (가짜 메일함)")
    r1 = p.run(target, force=True)
    print_result({k: (v.get("status") if isinstance(v, dict) and k != "collect" else v) for k, v in r1.items() if k != "collect"})
    c = r1.get("collect", {})
    print(f"→ 업체 {c.get('vendors')}곳 / {c.get('rows')}건 / {c.get('boxes')}박스 / 누락 {c.get('missing')}건")
    print("\n[나에게 온 요약 메일 본문]\n" + "-" * 70)
    print(mail.sent[-1][2])

    # 박경원님 회신: 운송장번호를 채운 엑셀
    supplier_mail = mail.sent[0]
    fname, data = supplier_mail[3][0]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["통합발주서"]
    for i in range(2, ws.max_row + 1):
        if i % 9:  # 몇 건은 미등록으로 남겨 '미등록' 안내 문구도 확인
            ws.cell(row=i, column=12).value = f"6071{i:08d}"
    buf = io.BytesIO(); wb.save(buf)
    mail.inject(cfg.supplier_name, cfg.supplier_email, "RE: " + supplier_mail[1], "운송장 등록했습니다. 첨부 확인하세요.",
                datetime.now(KST), [(fname, buf.getvalue())])

    print("\n" + "=" * 70, "\n[2단계] 박경원님 회신 감지 → 업체별 엑셀 + 안내문 발송")
    r2 = p.run(target)
    print_result(r2.get("reply"))
    vendor_mail = next(m for m in mail.sent if m[1].startswith("[영암부부농원]"))
    print("\n[업체에 발송되는 안내문 예시] 제목:", vendor_mail[1], "\n" + "-" * 70)
    print(vendor_mail[2])
    print("\n[나에게 온 2단계 요약]\n" + "-" * 70)
    print(mail.sent[-1][2])
    print("\n" + "=" * 70, "\n[3] 같은 날 다시 실행 → 중복 발송 없이 종료되는지")
    r3 = p.run(target)
    print_result({k: (v.get("status") if isinstance(v, dict) else v) for k, v in r3.items()})
    print(f"\n결과 파일: {demo_dir}  (.eml 은 실제로 발송될 메일 그대로이며 메일 프로그램으로 열어볼 수 있습니다)")
    for f in sorted(demo_dir.rglob("*")):
        if f.is_file():
            print("  -", f.relative_to(demo_dir))


def _synthetic_rows():
    import order_excel as oe
    names = ["김하늘", "이서준", "박지우", "최민서", "정도윤", "강수아", "조예준", "윤지호", "장하은", "임시우", "한유나", "오지안"]
    products = [("1개 500g 청무화과", 1)] * 6 + [("1개 1kg 청무화과", 1)] * 3 + [("1개 2kg 청무화과", 1), ("1개 1kg 홍청 세트", 2), ("1개 500g 청무화과", 3)]
    rows = []
    for i, (n, (item, q)) in enumerate(zip(names, products)):
        rows.append(oe.OrderRow(values={
            "받는분성함": n, "받는분우편번호": "", "받는분주소(전체, 분할)": f"서울특별시 테스트구 샘플로 {i + 1} 101동 {i + 1}01호" if i != 7 else "",
            "받는분전화번호": f"010-0000-{i:04d}", "받는분기타연락처": "", "박스수량": q, "배송메세지1": "문 앞",
            "품목명": item, "보내는분성명": "", "보내는분전화번호": "0507-0000-0000", "보내는분주소(전체, 분할)": "",
            "운송장번호": "", "원본품목(요약용)": oe.product_key(item)}))
    return rows


def main(argv=None):
    ap = argparse.ArgumentParser(description="영암부부농원 발주 자동화")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name, fn in [("run", cmd_run), ("collect", cmd_collect), ("check-reply", cmd_check_reply), ("daemon", cmd_daemon),
                     ("announce", cmd_announce), ("preview-announcement", cmd_preview_announcement),
                     ("test-login", cmd_test_login), ("demo", cmd_demo)]:
        sp = sub.add_parser(name)
        sp.set_defaults(fn=fn)
        sp.add_argument("--date", help="대상 날짜 (2026-10-05 또는 10.5)")
        sp.add_argument("--dry-run", action="store_true", help="메일을 실제로 보내지 않고 data/outbox 에 .eml 로 저장")
        sp.add_argument("--force", action="store_true", help="완료 표시가 있어도 다시 수행")
        sp.add_argument("--to", help="announce: 받는 업체 메일 (쉼표 구분)")
        sp.add_argument("--sample", help="demo: 업체 발주서로 쓸 샘플 엑셀 경로")
    args = ap.parse_args(argv)
    cfg = load_config()
    setup_logging(cfg)
    args.fn(cfg, args)


if __name__ == "__main__":
    main()
